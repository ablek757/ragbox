"""文档加载与文本切分。"""
from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Document:
    text: str
    source: str


def load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def load_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def load_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_url(url: str) -> str:
    import requests
    from bs4 import BeautifulSoup

    resp = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
    resp.raise_for_status()
    resp.encoding = resp.apparent_encoding or "utf-8"
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


EXTENSION_LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".md": load_text,
    ".txt": load_text,
    ".markdown": load_text,
}


def load_file(path: str | Path) -> Document:
    p = Path(path)
    ext = p.suffix.lower()
    loader = EXTENSION_LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"不支持的文件类型：{ext}")
    return Document(text=loader(p), source=p.name)


def load_from_url(url: str) -> Document:
    return Document(text=load_url(url), source=url)


_SPLIT_PATTERN = re.compile(r"(?<=[。！？!?\n])")


def split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    text = text.strip()
    if not text:
        return []

    pieces = [s for s in _SPLIT_PATTERN.split(text) if s.strip()]
    chunks: list[str] = []
    buf = ""
    for piece in pieces:
        if len(buf) + len(piece) <= chunk_size:
            buf += piece
        else:
            if buf:
                chunks.append(buf.strip())
            if len(piece) > chunk_size:
                for i in range(0, len(piece), chunk_size - overlap):
                    chunks.append(piece[i : i + chunk_size].strip())
                buf = ""
            else:
                buf = (chunks[-1][-overlap:] if chunks and overlap else "") + piece
    if buf.strip():
        chunks.append(buf.strip())
    return chunks


def chunk_documents(docs: Iterable[Document], chunk_size: int, overlap: int) -> list[dict]:
    out: list[dict] = []
    for doc in docs:
        for i, chunk in enumerate(split_text(doc.text, chunk_size, overlap)):
            out.append({"text": chunk, "source": doc.source, "chunk_id": i})
    return out
